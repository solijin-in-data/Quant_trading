import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as OpSide
from openpyxl.utils import get_column_letter
import os
import glob
import re
from datetime import datetime

# =========================================================================
# CẤU HÌNH THƯ MỤC HỆ THỐNG (BATCH PROCESSING)
# =========================================================================
INPUT_DIR = r"C:\Users\tuana\OneDrive\Documents\GitHub\Quant_trading\orderbookcheck\unprocess_data" 
OUTPUT_DIR = r"C:\Users\tuana\OneDrive\Documents\GitHub\Quant_trading\orderbookcheck\processed_data"
# =========================================================================

def parse_filename_info(filename):
    symbol_match = re.search(r'_([A-Z0-9]{3})_', filename)
    symbol = symbol_match.group(1) if symbol_match else "UNKNOWN"
    if symbol == "UNKNOWN":
        parts = filename.replace('.xlsx', '').split('_')
        if len(parts) >= 3:
            symbol = parts[-2]

    date_match = re.search(r'(\d{8})', filename)
    if date_match:
        date_str = date_match.group(1)
    else:
        date_str = datetime.now().strftime('%Y%m%d')
    return symbol, date_str

def process_single_file(file_path, symbol, date_str):
    filename = os.path.basename(file_path)
    output_filename = f"{symbol}_{date_str}_orderflow_phase1.xlsx"
    output_excel = os.path.join(OUTPUT_DIR, output_filename)
    
    print(f"\n" + "►"*80)
    print(f"🏗️  [ORDER FLOW PHASE 1] Phân tích: {symbol} | Ngày: {date_str}")
    
    # --- BƯỚC 1 & 2: DÒ TÌM HEADER VÀ LÀM SẠCH DỮ LIỆU ---
    df_temp = pd.read_excel(file_path, nrows=20, header=None)
    header_index = 0
    for i, row in df_temp.iterrows():
        row_str = ' '.join(row.astype(str).values).upper()
        if 'THỜI GIAN' in row_str or 'KIỂU' in row_str:
            header_index = i
            break
            
    df = pd.read_excel(file_path, skiprows=header_index)
    df.columns = [str(col).strip() for col in df.columns]
    
    rename_map = {'THỜI GIAN': 'Time', 'KIỂU': 'Side_Raw', 'GIÁ KHỚP': 'Price', 'THAY ĐỔI': 'Change', 'KL KHỚP': 'Volume'}
    df = df.rename(columns=rename_map)
    
    side_mapping = {'BU': 'BUY', 'SD': 'SELL', 'NA': 'ATC/ATO'}
    df['Side'] = df.get('Side_Raw', pd.Series(dtype='object')).map(side_mapping).fillna('ATC/ATO')
    
    for col in ['TỔNG KL KHỚP', '% THAY ĐỔI', 'Side_Raw']:
        if col in df.columns:
            df = df.drop(columns=[col])

    df['Price'] = pd.to_numeric(df['Price'], errors='coerce')
    df['Volume'] = pd.to_numeric(df['Volume'], errors='coerce')
    df = df.dropna(subset=['Price', 'Volume']).reset_index(drop=True)
    
    # Đảo ngược dòng để tính toán theo đúng trình tự thời gian (từ sáng đến chiều)
    df = df.iloc[::-1].reset_index(drop=True)

    # --- BƯỚC 3: ORDER FLOW ENGINE (PHASE 1) ---
    # 1. Net Volume & Cumulative Delta (CVD)
    df['Net_Volume'] = 0
    df.loc[df['Side'] == 'BUY', 'Net_Volume'] = df['Volume']
    df.loc[df['Side'] == 'SELL', 'Net_Volume'] = -df['Volume']
    df['CVD'] = df['Net_Volume'].cumsum()

    # 2. VWAP (Volume Weighted Average Price)
    df['Value'] = df['Price'] * df['Volume']
    df['CumValue'] = df['Value'].cumsum()
    df['CumVol'] = df['Volume'].cumsum()
    df['VWAP'] = df['CumValue'] / df['CumVol']

    # 3. Relative Volume (RVOL) & Whale Activity
    median_vol = df['Volume'].median()
    # Tránh chia cho 0 nếu median = 0
    df['RVOL'] = df['Volume'] / median_vol if median_vol > 0 else 0
    
    # Whale threshold: RVOL > 10 (Gấp 10 lần trung vị)
    df['Whale_Activity'] = 'NORMAL'
    df.loc[(df['RVOL'] >= 10) & (df['Side'] == 'BUY'), 'Whale_Activity'] = 'LARGE_BUY'
    df.loc[(df['RVOL'] >= 10) & (df['Side'] == 'SELL'), 'Whale_Activity'] = 'LARGE_SELL'

    # 4. Volume Profile & POC (Point of Control)
    vol_profile = df.groupby('Price')['Volume'].sum().sort_values(ascending=False)
    poc_price = vol_profile.index[0] if not vol_profile.empty else 0
    poc_vol = vol_profile.iloc[0] if not vol_profile.empty else 0

    # --- BƯỚC 4: BÁO CÁO TERMINAL ---
    print("\n" + "="*80)
    print(f"📊 REPORT PHASE 1: {symbol} - {date_str}")
    print("="*80)
    
    total_buy = df[df['Side'] == 'BUY']['Volume'].sum()
    total_sell = df[df['Side'] == 'SELL']['Volume'].sum()
    total_vol = df['Volume'].sum()
    
    print(f"• Tổng Volume: {total_vol:,.0f} CP")
    print(f"• VWAP Cuối ngày: {df['VWAP'].iloc[-1]:.2f}")
    print(f"• Cumulative Delta (CVD): {df['CVD'].iloc[-1]:,.0f} CP")
    print(f"• Median Volume (1 Tick): {median_vol:,.0f} CP")
    
    print("-" * 80)
    print("🎯 VOLUME PROFILE (TOP 3 VÙNG GIÁ GIAO TRANH MẠNH NHẤT - POC):")
    for idx, (price, vol) in enumerate(vol_profile.head(3).items(), 1):
        is_poc = " 👑 (POC)" if idx == 1 else ""
        print(f"   #{idx}: Giá {price:,.2f} -> Khớp {vol:,.0f} CP{is_poc}")
    print("="*80 + "\n")

    # --- BƯỚC 5: XUẤT EXCEL TỐI ƯU ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Order Flow {symbol}"
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["Thời Gian", "Giá Khớp", "Khối Lượng", "Bên", "RVOL", "Net Volume", "Cumulative Delta", "VWAP", "Radar Cá Mập"]
    ws.append(headers)
    
    for index, row in df.iterrows():
        time_val = row["Time"].strftime('%H:%M:%S') if hasattr(row["Time"], 'strftime') else row["Time"]
        ws.append([
            time_val, 
            row["Price"], 
            row["Volume"], 
            row["Side"], 
            round(row["RVOL"], 2), 
            row["Net_Volume"], 
            row["CVD"], 
            round(row["VWAP"], 2), 
            row["Whale_Activity"]
        ])
        
    # --- ĐỊNH DẠNG EXCEL ---
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(left=OpSide(style='thin', color='E0E0E0'), right=OpSide(style='thin', color='E0E0E0'), top=OpSide(style='thin', color='E0E0E0'), bottom=OpSide(style='thin', color='E0E0E0'))
    
    buy_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") 
    sell_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") 
    atc_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") 
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = center_align
        
    for r_idx in range(2, ws.max_row + 1):
        side_val = ws[f"D{r_idx}"].value
        ws.cell(row=r_idx, column=4).fill = buy_fill if side_val == "BUY" else sell_fill if side_val == "SELL" else atc_fill
        
        ws[f"A{r_idx}"].alignment = center_align
        ws[f"B{r_idx}"].alignment = right_align; ws[f"B{r_idx}"].number_format = '#,##0.00'
        ws[f"C{r_idx}"].alignment = right_align; ws[f"C{r_idx}"].number_format = '#,##0'
        ws[f"D{r_idx}"].alignment = center_align
        ws[f"E{r_idx}"].alignment = right_align; ws[f"E{r_idx}"].number_format = '#,##0.00'
        ws[f"F{r_idx}"].alignment = right_align; ws[f"F{r_idx}"].number_format = '#,##0'
        ws[f"G{r_idx}"].alignment = right_align; ws[f"G{r_idx}"].number_format = '#,##0'
        ws[f"H{r_idx}"].alignment = right_align; ws[f"H{r_idx}"].number_format = '#,##0.00'
        ws[f"I{r_idx}"].alignment = center_align
        
        for c_idx in range(1, len(headers) + 1):
            ws.cell(row=r_idx, column=c_idx).border = thin_border
            
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 16
        
    try:
        wb.save(output_excel)
        print(f"🎯 LƯU THÀNH CÔNG: {output_filename}")
    except PermissionError:
        print(f"❌ LỖI: Vui lòng đóng file {output_filename}!")

def run_batch_pipeline():
    if not os.path.exists(INPUT_DIR): os.makedirs(INPUT_DIR)
    if not os.path.exists(OUTPUT_DIR): os.makedirs(OUTPUT_DIR)
    excel_files = [f for f in glob.glob(os.path.join(INPUT_DIR, "*.xlsx")) if not os.path.basename(f).startswith("~$")]
    
    if not excel_files: return print(f"❌ Trống: {INPUT_DIR}")
    for file_path in excel_files:
        symbol, date_str = parse_filename_info(os.path.basename(file_path))
        process_single_file(file_path, symbol, date_str)

if __name__ == "__main__":
    run_batch_pipeline()