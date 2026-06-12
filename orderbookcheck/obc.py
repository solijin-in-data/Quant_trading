import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as OpSide
from openpyxl.utils import get_column_letter
import re
import os
from datetime import datetime

# =========================================================================
# CẤU HÌNH THÔNG TIN
# =========================================================================
SYMBOL = "CTD"
FOLDER_PATH = r"C:\Users\tuana\OneDrive\Documents\GitHub\Quant_trading\orderbookcheck\data"
INPUT_FILE_NAME = "CTD_6_12_2026.txt" # Tên file data bạn đã cào được
# =========================================================================

def parse_and_analyze():
    file_path = os.path.join(FOLDER_PATH, INPUT_FILE_NAME)
    print(f"=== BƯỚC 1: ĐANG ĐỌC VÀ BÓC TÁCH DỮ LIỆU TỪ: {INPUT_FILE_NAME} ===")
    
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            html_content = f.read()
    except FileNotFoundError:
        print(f"❌ KHÔNG TÌM THẤY FILE: {file_path}")
        return None

    # Thuật toán quét HTML
    pattern = r'class="list-row[^"]*".*?<div style="display: flex; flex-direction: row; flex: 1 1 0%;">(.*?)<\/div><\/div><\/div>'
    matches = re.findall(pattern, html_content, re.DOTALL)
    if not matches:
        pattern = r'<div style="display: flex; flex-direction: row; flex: 1 1 0%;">(.*?)<\/div><\/div><\/div>'
        matches = re.findall(pattern, html_content, re.DOTALL)
    
    records = []
    for block in matches:
        divs = re.findall(r'<div[^>]*color:\s*rgb\(([^)]+)\)[^>]*>([^<]+)</div>|<div[^>]*>([^<]+)</div>', block)
        parsed_divs = []
        for match in divs:
            color = match[0] if match[0] else ""
            text = match[1].strip() if match[1] else match[2].strip()
            if text:
                parsed_divs.append({"text": text, "color": color})
                
        if len(parsed_divs) >= 4:
            time_val = parsed_divs[0]["text"]
            if re.match(r'^\d{2}:\d{2}:\d{2}$', time_val):
                price_val = parsed_divs[1]["text"]
                change_val = parsed_divs[2]["text"]
                volume_val = parsed_divs[3]["text"]
                price_color = parsed_divs[1]["color"]
                side_raw = parsed_divs[4]["text"] if len(parsed_divs) >= 5 else ""
                
                if '0, 170, 0' in price_color or side_raw in ['S', 'M']:
                    side = 'BUY'
                elif '238, 84, 66' in price_color or side_raw == 'B':
                    side = 'SELL'
                else:
                    side = 'ATC/ATO'
                    
                records.append({
                    "Time": time_val, "Price": price_val, 
                    "Change": change_val, "Volume": volume_val, "Side": side
                })

    df = pd.DataFrame(records)
    if df.empty:
        print("❌ Cấu trúc HTML trống.")
        return None
        
    df = df.drop_duplicates(subset=["Time", "Price", "Volume", "Side"]).reset_index(drop=True)
    df = df.iloc[::-1].reset_index(drop=True) 
    
    df["Price"] = df["Price"].astype(str).str.replace(",", "", regex=False)
    df["Price"] = pd.to_numeric(df["Price"], errors='coerce')
    df["Volume"] = df["Volume"].astype(str).str.replace(",", "", regex=False).str.replace(".", "", regex=False)
    df["Volume"] = pd.to_numeric(df["Volume"], errors='coerce').fillna(0).astype(int)
    df = df.dropna(subset=["Price"]).reset_index(drop=True)
    
    print(f"✅ Đã bóc tách thành công {len(df)} lệnh khớp sạch của mã {SYMBOL}!")

    # =========================================================================
    # PHẦN MỚI: BOT PHÂN TÍCH ĐỊNH LƯỢNG (QUANTITATIVE ANALYSIS)
    # =========================================================================
    print("\n" + "="*60)
    print(f"📊 BÁO CÁO PHÂN TÍCH DÒNG TIỀN MÃ {SYMBOL} 📊")
    print("="*60)

    # 1. Thống kê Áp lực Mua / Bán
    total_buy = df[df['Side'] == 'BUY']['Volume'].sum()
    total_sell = df[df['Side'] == 'SELL']['Volume'].sum()
    total_atc = df[df['Side'] == 'ATC/ATO']['Volume'].sum()
    total_vol = df['Volume'].sum()

    buy_pct = (total_buy / total_vol) * 100 if total_vol > 0 else 0
    sell_pct = (total_sell / total_vol) * 100 if total_vol > 0 else 0

    print(f"🔹 1. TỔNG QUAN CUNG CẦU CHỦ ĐỘNG:")
    print(f"   - Tổng khối lượng khớp: {total_vol:,.0f} CP")
    print(f"   - 🟢 MUA chủ động:      {total_buy:,.0f} CP ({buy_pct:.1f}%)")
    print(f"   - 🔴 BÁN chủ động:      {total_sell:,.0f} CP ({sell_pct:.1f}%)")
    
    net_vol = total_buy - total_sell
    if net_vol > 0:
        print(f"   => KẾT LUẬN: Phe MUA đang áp đảo, bơm ròng +{net_vol:,.0f} CP vào thị trường.")
    elif net_vol < 0:
        print(f"   => KẾT LUẬN: Phe BÁN đang xả hàng, rút ròng {net_vol:,.0f} CP khỏi thị trường.")
    else:
        print(f"   => KẾT LUẬN: Cung cầu đang giằng co cân bằng.")

    # 2. Phân tích vùng giá tập trung (Point of Control)
    print(f"\n🔹 2. HỒ SƠ KHỐI LƯỢNG (Top 3 Vùng giá giao dịch nhiều nhất):")
    vol_by_price = df.groupby('Price')['Volume'].sum().sort_values(ascending=False).head(3)
    for price, vol in vol_by_price.items():
        print(f"   - Mức giá {price}: {vol:,.0f} CP")

    # 3. Radar Quét Cá Mập (Lấy Top 5% lệnh to nhất hoặc các lệnh trên ngưỡng cố định)
    # Thuật toán tính ngưỡng: Lấy phân vị 95% (95th percentile)
    whale_threshold = df['Volume'].quantile(0.95)
    whales = df[df['Volume'] >= whale_threshold]
    
    print(f"\n🔹 3. RADAR CÁ MẬP (Lọc các lệnh Đột biến >= {whale_threshold:,.0f} CP):")
    if not whales.empty:
        print(f"   🚨 CẢNH BÁO: Phát hiện {len(whales)} lệnh siêu lớn!")
        # Sắp xếp để xem lệnh to nhất
        whales_sorted = whales.sort_values(by='Volume', ascending=False).head(10) 
        for idx, row in whales_sorted.iterrows():
            icon = "🟢 MUA" if row['Side'] == 'BUY' else "🔴 BÁN" if row['Side'] == 'SELL' else "⚪ ATC/ATO"
            print(f"      [{row['Time']}] {icon} tấp tay {row['Volume']:,.0f} CP ở giá {row['Price']}")
        if len(whales) > 10:
            print(f"      ... (và {len(whales) - 10} lệnh cá mập khác được lưu trong Excel)")
    else:
        print("   - Không phát hiện lệnh cá mập bất thường. Thị trường chủ yếu là lệnh nhỏ lẻ.")
    
    print("="*60 + "\n")

    # =========================================================================
    print("=== BƯỚC 3: TỰ ĐỘNG KHỞI TẠO VÀ FORMAT FILE EXCEL ===")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Sổ Lệnh {SYMBOL}"
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["Thời Gian", "Giá Khớp", "Biến Động", "Khối Lượng", "Bên Chủ Động", "Giá Trị Khớp (VND)"]
    ws.append(headers)
    
    for index, row in df.iterrows():
        ws.append([row["Time"], row["Price"], row["Change"], row["Volume"], row["Side"], ""])
        r_num = index + 2
        ws[f"F{r_num}"] = f"=B{r_num}*1000*D{r_num}"
        
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=OpSide(style='thin', color='D9D9D9'), right=OpSide(style='thin', color='D9D9D9'),
        top=OpSide(style='thin', color='D9D9D9'), bottom=OpSide(style='thin', color='D9D9D9')
    )
    
    buy_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  
    sell_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") 
    atc_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  
    
    for col_num in range(1, 7):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for r_idx in range(2, ws.max_row + 1):
        side_val = ws[f"E{r_idx}"].value
        current_fill = buy_fill if side_val == "BUY" else sell_fill if side_val == "SELL" else atc_fill
        
        ws[f"A{r_idx}"].alignment = center_align
        ws[f"B{r_idx}"].alignment = right_align
        ws[f"B{r_idx}"].number_format = '#,##0.00'
        ws[f"C{r_idx}"].alignment = right_align
        ws[f"C{r_idx}"].number_format = '#,##0.00'
        ws[f"D{r_idx}"].alignment = right_align
        ws[f"D{r_idx}"].number_format = '#,##0'
        ws[f"E{r_idx}"].alignment = center_align
        ws[f"F{r_idx}"].alignment = right_align
        ws[f"F{r_idx}"].number_format = '#,##0'
        
        for c_idx in range(1, 7):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = current_fill
            cell.border = thin_border
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    # Thêm HHMMSS vào tên file để không bao giờ bị lỗi ghi đè (Permission Denied) nếu đang mở Excel
    date_str = datetime.now().strftime("%d_%m_%Y_%H%M%S")
    output_excel = os.path.join(FOLDER_PATH, f"{SYMBOL}_{date_str}.xlsx")
    
    wb.save(output_excel)
    print(f"🚀 THÀNH CÔNG RỰC RỠ! ĐÃ XUẤT FILE EXCEL TẠI: {output_excel}")
    return df

if __name__ == "__main__":
    parse_and_analyze()