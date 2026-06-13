import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as OpSide
from openpyxl.utils import get_column_letter
import os

# =========================================================================
# CẤU HÌNH HỆ THỐNG TỪ DỮ LIỆU KAFI (EXCEL GỐC)
# =========================================================================
SYMBOL = "ACB"  
FOLDER_PATH = r"D:\Github\Quant_trading\orderbookcheck\data"
INPUT_FILE_NAME = "Kafi_Time&Sales_ACB_20260613.xlsx" # Đọc thẳng file Excel
OUTPUT_FILE_NAME = f"{SYMBOL}_orderbook.xlsx"
# =========================================================================

def run_kafi_excel_pipeline():
    file_path = os.path.join(FOLDER_PATH, INPUT_FILE_NAME)
    output_excel = os.path.join(FOLDER_PATH, OUTPUT_FILE_NAME)
    
    print(f"🏗️  [PIPELINE] Bắt đầu đọc trực tiếp Excel Kafi mã {SYMBOL}...")
    
    if not os.path.exists(file_path):
        print(f"❌ THẤT BẠI: Không tìm thấy file dữ liệu tại: {file_path}")
        return None

    # --- BƯỚC 1: DÒ TÌM THÔNG MINH DÒNG CHỨA HEADER TRONG EXCEL ---
    print("🔍 Đang dò tìm dòng tiêu đề chuẩn để loại bỏ thông tin rác...")
    # Đọc nháp 20 dòng đầu tiên, không lấy header
    df_temp = pd.read_excel(file_path, nrows=20, header=None)
    header_index = 0
    
    for i, row in df_temp.iterrows():
        # Gộp toàn bộ text trong dòng thành 1 chuỗi để tìm kiếm
        row_str = ' '.join(row.astype(str).values).upper()
        if 'THỜI GIAN' in row_str or 'KIỂU' in row_str:
            header_index = i
            break
            
    # --- BƯỚC 2: ĐỌC VÀ LÀM SẠCH DỮ LIỆU ---
    # Nạp lại file Excel và bỏ qua chính xác số dòng rác bên trên
    df = pd.read_excel(file_path, skiprows=header_index)
    
    df.columns = [str(col).strip() for col in df.columns]
    
    rename_map = {
        'THỜI GIAN': 'Time',
        'KIỂU': 'Side_Raw',
        'GIÁ KHỚP': 'Price',
        'THAY ĐỔI': 'Change',
        'KL KHỚP': 'Volume'
    }
    df = df.rename(columns=rename_map)
    
    # Chuyển đổi mã lệnh Kafi
    side_mapping = {'BU': 'BUY', 'SD': 'SELL', 'NA': 'ATC/ATO'}
    df['Side'] = df.get('Side_Raw', pd.Series(dtype='object')).map(side_mapping).fillna('ATC/ATO')
    
    # Dọn dẹp cột thừa
    for col in ['TỔNG KL KHỚP', '% THAY ĐỔI', 'Side_Raw']:
        if col in df.columns:
            df = df.drop(columns=[col])

    # Ép kiểu dữ liệu số và xóa dòng rỗng
    df['Price'] = pd.to_numeric(df['Price'], errors='coerce')
    df['Volume'] = pd.to_numeric(df['Volume'], errors='coerce')
    df = df.dropna(subset=['Price', 'Volume']).reset_index(drop=True)
    
    # Đảo trục thời gian
    df = df.iloc[::-1].reset_index(drop=True)
    
    print(f"✅ Đọc và chuẩn hóa thành công {len(df)} dòng lệnh giao dịch (Tick Data).")

    # --- BƯỚC 3: ENGINE TÍNH TOÁN CẤU TRÚC VI MÔ ---
    print("🧠 Đang chạy Engine phân tích Order Flow...")
    
    df['Net_Volume'] = 0
    df.loc[df['Side'] == 'BUY', 'Net_Volume'] = df['Volume']
    df.loc[df['Side'] == 'SELL', 'Net_Volume'] = -df['Volume']
    
    window = 5
    df['Vol_Roll_Sum'] = df['Volume'].rolling(window=window).sum()
    df['Price_Range'] = df['Price'].rolling(window=window).apply(lambda x: max(x) - min(x) + 0.01, raw=True)
    df['Absorption_Score'] = df['Vol_Roll_Sum'] / df['Price_Range']
    
    # Radar Cá mập (Phân vị 95%)
    whale_threshold = df['Volume'].quantile(0.95)
    df['Whale_Activity'] = 'NORMAL'
    df.loc[(df['Volume'] >= whale_threshold) & (df['Side'] == 'BUY'), 'Whale_Activity'] = 'LARGE_BUY'
    df.loc[(df['Volume'] >= whale_threshold) & (df['Side'] == 'SELL'), 'Whale_Activity'] = 'LARGE_SELL'

    # --- BƯỚC 4: IN BÁO CÁO PHÂN TÍCH LÊN CMD ---
    print("\n" + "="*80)
    print(f"📊 THỐNG KÊ DÒNG TIỀN PHIÊN (MÃ {SYMBOL} - DATA: KAFI):")
    print("="*80)
    
    total_buy = df[df['Side'] == 'BUY']['Volume'].sum()
    total_sell = df[df['Side'] == 'SELL']['Volume'].sum()
    total_vol = df['Volume'].sum()
    
    buy_pct = (total_buy / total_vol) * 100 if total_vol > 0 else 0
    sell_pct = (total_sell / total_vol) * 100 if total_vol > 0 else 0
    
    print(f"• Tổng khối lượng khớp: {total_vol:,.0f} CP")
    print(f"• 🟢 MUA chủ động (BU):    {total_buy:,.0f} CP ({buy_pct:.1f}%)")
    print(f"• 🔴 BÁN chủ động (SD):    {total_sell:,.0f} CP ({sell_pct:.1f}%)")
    
    net_flow = total_buy - total_sell
    if net_flow > 0:
        print(f"• Net Order Flow Imbalance: +{net_flow:,.0f} cp (Phe MUA ăn lên mạnh mẽ)")
    else:
        print(f"• Net Order Flow Imbalance: {net_flow:,.0f} cp (Phe BÁN xả hàng áp đảo)")
        
    print("-" * 80)
    print("🚨 TOP 5 VÙNG GIÁ CÁ MẬP ĐẶT LỆNH ẨN HẤP THỤ MẠNH NHẤT:")
    top_absorption = df.sort_values(by='Absorption_Score', ascending=False).head(5)
    print(top_absorption[['Time', 'Price', 'Volume', 'Side', 'Absorption_Score']].to_string(index=False))
    print("="*80 + "\n")

    # --- BƯỚC 5: XUẤT EXCEL CHUẨN HÓA ---
    print(f"💾 Đang dựng file Excel: {OUTPUT_FILE_NAME}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Sổ Lệnh {SYMBOL}"
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["Thời Gian", "Giá Khớp", "Biến Động", "Khối Lượng", "Bên Chủ Động", "Giá Trị Khớp (VND)", "Dòng Tiền Ròng", "Điểm Hấp Thụ", "Radar Cá Mập"]
    ws.append(headers)
    
    # Định dạng thời gian chuẩn để tránh hiển thị dạng datetime object dài dòng
    for index, row in df.iterrows():
        # Xử lý format Time nếu là object datetime.time của pandas/excel
        time_val = row["Time"]
        if hasattr(time_val, 'strftime'):
            time_val = time_val.strftime('%H:%M:%S')
            
        ws.append([time_val, row["Price"], row["Change"], row["Volume"], row["Side"], "", row["Net_Volume"], row["Absorption_Score"], row["Whale_Activity"]])
        r_num = index + 2
        ws[f"F{r_num}"] = f"=B{r_num}*1000*D{r_num}"
        
    # --- ĐỊNH DẠNG MÀU SẮC EXCEL ---
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(left=OpSide(style='thin', color='E0E0E0'), right=OpSide(style='thin', color='E0E0E0'), top=OpSide(style='thin', color='E0E0E0'), bottom=OpSide(style='thin', color='E0E0E0'))
    
    buy_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") 
    sell_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") 
    atc_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") 
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for r_idx in range(2, ws.max_row + 1):
        side_val = ws[f"E{r_idx}"].value
        current_fill = buy_fill if side_val == "BUY" else sell_fill if side_val == "SELL" else atc_fill
        
        ws[f"A{r_idx}"].alignment = center_align
        ws[f"B{r_idx}"].alignment = right_align; ws[f"B{r_idx}"].number_format = '#,##0.00'
        ws[f"C{r_idx}"].alignment = right_align; ws[f"C{r_idx}"].number_format = '#,##0.00'
        ws[f"D{r_idx}"].alignment = right_align; ws[f"D{r_idx}"].number_format = '#,##0'
        ws[f"E{r_idx}"].alignment = center_align
        ws[f"F{r_idx}"].alignment = right_align; ws[f"F{r_idx}"].number_format = '#,##0'
        ws[f"G{r_idx}"].alignment = right_align; ws[f"G{r_idx}"].number_format = '#,##0'
        ws[f"H{r_idx}"].alignment = right_align; ws[f"H{r_idx}"].number_format = '#,##0.00'
        ws[f"I{r_idx}"].alignment = center_align
        
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = current_fill
            cell.border = thin_border
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    try:
        wb.save(output_excel)
        print(f"🎯 XUẤT SẮC: Đã xuất file chuẩn tại: {output_excel}")
    except PermissionError:
        print(f"❌ LỖI: Vui lòng đóng file {OUTPUT_FILE_NAME} đang mở trong Excel!")

if __name__ == "__main__":
    run_kafi_excel_pipeline()