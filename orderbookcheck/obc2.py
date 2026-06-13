import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as OpSide
from openpyxl.utils import get_column_letter
import re
import os

# =========================================================================
# CẤU HÌNH THÔNG TIN HỆ THỐNG QUANT
# =========================================================================
SYMBOL = "CTD"
FOLDER_PATH = r"D:\Github\Quant_trading\orderbookcheck\data"
INPUT_FILE_NAME = "CTD_Full_Data_Manual.txt"   # File dữ liệu HTML thô của bạn
OUTPUT_FILE_NAME = "CTD_orderbook.xlsx"        # Tên file Excel chuẩn hóa cố định
# =========================================================================

def run_quant_pipeline():
    file_path = os.path.join(FOLDER_PATH, INPUT_FILE_NAME)
    output_excel = os.path.join(FOLDER_PATH, OUTPUT_FILE_NAME)
    
    print(f"🏗️  [PIPELINE] Bắt đầu xử lý dữ liệu cấu trúc vi mô mã {SYMBOL}...")
    print(f"📂 Đang đọc file HTML thô: {INPUT_FILE_NAME}")
    
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            html_content = f.read()
    except FileNotFoundError:
        print(f"❌ THẤT BẠI: Không tìm thấy file dữ liệu tại: {file_path}")
        return None

    # --- BƯỚC 1: BÓC TÁCH MÃ HTML ---
    pattern = r'class="list-row[^"]*".*?<div style="display: flex; flex-direction: row; flex: 1 1 0%;">(.*?)<\/div><\/div><\/div>'
    matches = re.findall(pattern, html_content, re.DOTALL)
    if not matches:
        pattern = r'<div style="display: flex; flex-direction: row; flex: 1 1 0%;">(.*?)<\/div><\/div><\/div>'
        matches = re.findall(pattern, html_content, re.DOTALL)
    
    records = []
    for block in matches:
        divs = re.findall(r'<div[^>]*color:\s*rgb\(([^)]+)\)[^>]*>([^<]+)</div>|<div[^>]*>([^<]+)</div>', block)
        parsed_divs = []
        for match in divs:
            color = match[0] if match[0] else ""
            text = match[1].strip() if match[1] else match[2].strip()
            if text:
                parsed_divs.append({"text": text, "color": color})
                
        if len(parsed_divs) >= 4:
            time_val = parsed_divs[0]["text"]
            if re.match(r'^\d{2}:\d{2}:\d{2}$', time_val):
                price_val = parsed_divs[1]["text"]
                change_val = parsed_divs[2]["text"]
                volume_val = parsed_divs[3]["text"]
                price_color = parsed_divs[1]["color"]
                side_raw = parsed_divs[4]["text"] if len(parsed_divs) >= 5 else ""
                
                if '0, 170, 0' in price_color or side_raw in ['M', 'BUY']:
                    side = 'BUY'
                elif '238, 84, 66' in price_color or side_raw in ['B', 'SELL']:
                    side = 'SELL'
                else:
                    side = 'ATC/ATO'
                    
                records.append({
                    "Time": time_val.strip(), "Price": price_val, 
                    "Change": change_val, "Volume": volume_val, "Side": side
                })

    df = pd.DataFrame(records)
    if df.empty:
        print("❌ LỖI: Cấu trúc HTML trống hoặc không khớp mẫu quét RegEx.")
        return None
        
    # --- BƯỚC 2: LÀM SẠCH DỮ LIỆU SỐ ---
    df = df.drop_duplicates(subset=["Time", "Price", "Volume", "Side"]).reset_index(drop=True)
    df = df.iloc[::-1].reset_index(drop=True) 
    
    df["Price"] = df["Price"].astype(str).str.replace(",", "", regex=False)
    df["Price"] = pd.to_numeric(df["Price"], errors='coerce')
    df["Volume"] = df["Volume"].astype(str).str.replace(",", "", regex=False).str.replace(".", "", regex=False)
    df["Volume"] = pd.to_numeric(df["Volume"], errors='coerce').fillna(0).astype(int)
    df = df.dropna(subset=["Price"]).reset_index(drop=True)
    
    print(f"✅ Bóc tách thành công {len(df)} dòng lệnh giao dịch (Tick Data).")

    # --- BƯỚC 3: TÍNH TOÁN ĐẶC TRƯNG CẤU TRÚC VI MÔ ---
    print("🧠 Đang tính toán chỉ báo Order Flow: Net Volume & Absorption Score...")
    
    df['Net_Volume'] = 0
    df.loc[df['Side'] == 'BUY', 'Net_Volume'] = df['Volume']
    df.loc[df['Side'] == 'SELL', 'Net_Volume'] = -df['Volume']
    
    window = 5
    df['Vol_Roll_Sum'] = df['Volume'].rolling(window=window).sum()
    df['Price_Range'] = df['Price'].rolling(window=window).apply(lambda x: max(x) - min(x) + 0.01, raw=True)
    df['Absorption_Score'] = df['Vol_Roll_Sum'] / df['Price_Range']
    
    whale_threshold = df['Volume'].quantile(0.95)
    df['Whale_Activity'] = 'NORMAL'
    df.loc[(df['Volume'] >= whale_threshold) & (df['Side'] == 'BUY'), 'Whale_Activity'] = 'LARGE_BUY'
    df.loc[(df['Volume'] >= whale_threshold) & (df['Side'] == 'SELL'), 'Whale_Activity'] = 'LARGE_SELL'

    # =========================================================================
    # PHẦN ĐƯỢC KHÔI PHỤC VÀ NÂNG CẤP: BẢN IN PHÂN TÍCH RA CMD
    # =========================================================================
    print("\n" + "="*80)
    print(f"📊 THỐNG KÊ DÒNG TIỀN PHIÊN 12/06/2026 (MÃ {SYMBOL}):")
    print("="*80)
    
    total_buy = df[df['Side'] == 'BUY']['Volume'].sum()
    total_sell = df[df['Side'] == 'SELL']['Volume'].sum()
    total_atc = df[df['Side'] == 'ATC/ATO']['Volume'].sum()
    total_vol = df['Volume'].sum()
    
    buy_pct = (total_buy / total_vol) * 100 if total_vol > 0 else 0
    sell_pct = (total_sell / total_vol) * 100 if total_vol > 0 else 0
    
    print(f"• Tổng khối lượng khớp: {total_vol:,.0f} CP")
    print(f"• 🟢 MUA chủ động:       {total_buy:,.0f} CP ({buy_pct:.1f}%)")
    print(f"• 🔴 BÁN chủ động:       {total_sell:,.0f} CP ({sell_pct:.1f}%)")
    
    net_flow = total_buy - total_sell
    if net_flow > 0:
        print(f"• Net Order Flow Imbalance: +{net_flow:,} cp (Bên Mua áp đảo)")
    elif net_flow < 0:
        print(f"• Net Order Flow Imbalance: {net_flow:,} cp (Bên Bán đè bẹp)")
    else:
        print(f"• Net Order Flow Imbalance: 0 cp (Cung cầu cân bằng)")
        
    print("-" * 80)
    print("🚨 VÙNG PHÁT HIỆN HẤP THỤ LỆNH (ABSORPTION) MẠNH NHẤT PHIÊN:")
    
    # Sắp xếp lấy 5 vùng có điểm hấp thụ cao nhất
    top_absorption = df.sort_values(by='Absorption_Score', ascending=False).head(5)
    print(top_absorption[['Time', 'Price', 'Volume', 'Side', 'Absorption_Score']].to_string(index=False))
    print("="*80 + "\n")

    # --- BƯỚC 4: KHỞI TẠO VÀ ĐỊNH DẠNG EXCEL ---
    print(f"💾 Đang dựng file Excel cấu trúc: {OUTPUT_FILE_NAME}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Sổ Lệnh {SYMBOL}"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "Thời Gian", "Giá Khớp", "Biến Động", "Khối Lượng", 
        "Bên Chủ Động", "Giá Trị Khớp (VND)", "Dòng Tiền Ròng", "Điểm Hấp Thụ", "Radar Cá Mập"
    ]
    ws.append(headers)
    
    for index, row in df.iterrows():
        ws.append([
            row["Time"], row["Price"], row["Change"], row["Volume"], 
            row["Side"], "", row["Net_Volume"], row["Absorption_Score"], row["Whale_Activity"]
        ])
        r_num = index + 2
        ws[f"F{r_num}"] = f"=B{r_num}*1000*D{r_num}"
        
    # Formatting hình ảnh Excel
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=OpSide(style='thin', color='E0E0E0'), right=OpSide(style='thin', color='E0E0E0'),
        top=OpSide(style='thin', color='E0E0E0'), bottom=OpSide(style='thin', color='E0E0E0')
    )
    buy_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") 
    sell_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") 
    atc_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") 
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for r_idx in range(2, ws.max_row + 1):
        side_val = ws[f"E{r_idx}"].value
        current_fill = buy_fill if side_val == "BUY" else sell_fill if side_val == "SELL" else atc_fill
        
        ws[f"A{r_idx}"].alignment = center_align
        ws[f"B{r_idx}"].alignment = right_align
        ws[f"B{r_idx}"].number_format = '#,##0.00'
        ws[f"C{r_idx}"].alignment = right_align
        ws[f"C{r_idx}"].number_format = '#,##0.00'
        ws[f"D{r_idx}"].alignment = right_align
        ws[f"D{r_idx}"].number_format = '#,##0'
        ws[f"E{r_idx}"].alignment = center_align
        ws[f"F{r_idx}"].alignment = right_align
        ws[f"F{r_idx}"].number_format = '#,##0'
        ws[f"G{r_idx}"].alignment = right_align
        ws[f"G{r_idx}"].number_format = '#,##0'
        ws[f"H{r_idx}"].alignment = right_align
        ws[f"H{r_idx}"].number_format = '#,##0.00'
        ws[f"I{r_idx}"].alignment = center_align
        
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = current_fill
            cell.border = thin_border
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    try:
        wb.save(output_excel)
        print(f"🎯 PIPELINE HOÀN THÀNH: Đã xuất file cấu trúc vi mô cố định tại: {output_excel}")
    except PermissionError:
        print(f"❌ LỖI GHI FILE: Hãy tắt file {OUTPUT_FILE_NAME} đang mở trong Excel!")

if __name__ == "__main__":
    run_quant_pipeline()